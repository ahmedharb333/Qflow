import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

wb = openpyxl.Workbook()
FONT = "Arial"
NAVY = "1F3864"
LIGHT = "DCE6F1"
INPUT_FILL = "FFFDE0"
thin = Side(style="thin", color="BFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

def label(ws, cell_ref, text, bold=True, fill=NAVY, color="FFFFFF", size=10):
    c = ws[cell_ref]
    c.value = text
    c.font = Font(name=FONT, size=size, bold=bold, color=color)
    if fill:
        c.fill = PatternFill("solid", fgColor=fill)
    c.alignment = Alignment(vertical="center")

def input_cell(ws, rng, size=10, align="left"):
    ws.merge_cells(rng)
    c = ws[rng.split(":")[0]]
    c.font = Font(name=FONT, size=size)
    c.fill = PatternFill("solid", fgColor=INPUT_FILL)
    c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=True)
    for row in ws[rng]:
        for cell in row:
            cell.border = border

# ---------- Sheet 1: Process Template ----------
ws = wb.active
ws.title = "Process Template"
ws.sheet_view.showGridLines = False
for i, w in enumerate([4, 20, 16, 20, 16, 4], start=1):
    ws.column_dimensions[get_column_letter(i)].width = w

ws.merge_cells("B2:E2")
ws["B2"] = "QFLOW CONSULTING — PROCESS TEMPLATE"
ws["B2"].font = Font(name=FONT, size=16, bold=True, color=NAVY)

fields = [
    ("Process Name", "[e.g. Internal Audit Process]"),
    ("Process Owner", "[Role/Name]"),
    ("Process No.", "[PRC-XXX]"),
    ("Version / Date", "[v1.0 — DD-MMM-YYYY]"),
    ("Objective", "[Why this process exists — the outcome it delivers]"),
    ("Scope", "[Where it starts, where it ends, what/who it covers]"),
    ("Trigger / Start Event", "[What kicks the process off]"),
    ("End Result", "[What marks the process complete]"),
]
r = 4
for name, placeholder in fields:
    label(ws, f"B{r}", name, fill=LIGHT, color="000000")
    input_cell(ws, f"C{r}:E{r}", align="left")
    ws[f"C{r}"] = placeholder
    r += 1

r += 1
label(ws, f"B{r}", "INPUTS", fill=NAVY)
label(ws, f"C{r}", "OUTPUTS", fill=NAVY)
label(ws, f"D{r}", "RELATED DOCUMENTS", fill=NAVY)
label(ws, f"E{r}", "APPLICABLE ISO CLAUSE(S)", fill=NAVY)
r += 1
for _ in range(4):
    for col in ["B", "C", "D", "E"]:
        input_cell(ws, f"{col}{r}:{col}{r}")
    r += 1

r += 1
label(ws, f"B{r}", "PROCESS STEPS", fill=NAVY)
ws.merge_cells(f"B{r}:E{r}")
r += 1
headers = ["Step No.", "Step Description", "Responsible", "Records / Evidence"]
widths_map = ["B", "C", "D", "E"]
for col, h in zip(widths_map, headers):
    label(ws, f"{col}{r}", h, fill=LIGHT, color="000000")
step_header_row = r
r += 1
first_step_row = r
sample = [
    (1, "[e.g. Schedule internal audit]", "[Quality Manager]", "[Audit schedule]"),
    (2, "", "", ""),
    (3, "", "", ""),
    (4, "", "", ""),
    (5, "", "", ""),
    (6, "", "", ""),
]
for num, desc, resp, rec in sample:
    ws[f"B{r}"] = num
    ws[f"C{r}"] = desc
    ws[f"D{r}"] = resp
    ws[f"E{r}"] = rec
    for col in ["B", "C", "D", "E"]:
        c = ws[f"{col}{r}"]
        c.border = border
        c.font = Font(name=FONT, size=10)
        c.alignment = Alignment(wrap_text=True, vertical="center")
        if col != "B":
            c.fill = PatternFill("solid", fgColor=INPUT_FILL)
    r += 1
last_step_row = r - 1

r += 1
label(ws, f"B{r}", "REVISION HISTORY", fill=NAVY)
ws.merge_cells(f"B{r}:E{r}")
r += 1
for col, h in zip(["B", "C", "D", "E"], ["Version", "Date", "Description of Change", "Approved By"]):
    label(ws, f"{col}{r}", h, fill=LIGHT, color="000000")
r += 1
for _ in range(3):
    for col in ["B", "C", "D", "E"]:
        input_cell(ws, f"{col}{r}:{col}{r}")
    r += 1

r += 1
ws.merge_cells(f"B{r}:E{r}")
ws[f"B{r}"] = "Yellow cells = fields to edit. For a visual process map, use Insert > SmartArt/Drawing (Word/Excel) or paste steps into the Procedure Template's swimlane generator."
ws[f"B{r}"].font = Font(name=FONT, size=9, italic=True, color="808080")

wb.save("Qflow_Process_Template.xlsx")
print("written")

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Invoice"

FONT = "Arial"
NAVY = "1F3864"
LIGHT = "DCE6F1"
YELLOW = "FFFF00"
INPUT_FILL = "FFFDE0"

thin = Side(style="thin", color="BFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

def set_col_widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

set_col_widths(ws, [4, 22, 12, 12, 14, 14, 4])

def merge_val(rng, text, size=11, bold=False, color="000000", align="left", fill=None, italic=False):
    ws.merge_cells(rng)
    c = ws[rng.split(":")[0]]
    c.value = text
    c.font = Font(name=FONT, size=size, bold=bold, color=color, italic=italic)
    c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=True)
    if fill:
        for row in ws[rng]:
            for cell in row:
                cell.fill = PatternFill("solid", fgColor=fill)

# Header / letterhead
merge_val("B2:D2", "QFLOW CONSULTING", size=18, bold=True, color=NAVY)
merge_val("B3:D3", "ISO & Quality Management Consulting", size=10, color="595959")
merge_val("B4:D4", "Abu Dhabi, United Arab Emirates | TRN: [•]", size=10, color="595959")
merge_val("B5:D5", "www.qflowconsulting.com | [email]", size=10, color="595959")

merge_val("E2:F2", "INVOICE", size=20, bold=True, color=NAVY, align="right")
merge_val("E3:F3", "Invoice No.: [•]", size=10, align="right")
merge_val("E4:F4", "Invoice Date: [DD-MMM-YYYY]", size=10, align="right")
merge_val("E5:F5", "Due Date: [DD-MMM-YYYY]", size=10, align="right")
merge_val("E6:F6", "Status: DRAFT", size=10, bold=True, align="right", color=NAVY)

# Bill to
merge_val("B7:C7", "BILL TO", size=10, bold=True, color="FFFFFF", fill=NAVY)
merge_val("B8:C8", "[Client Legal Name]")
merge_val("B9:C9", "[Client Address]")
merge_val("B10:C10", "TRN: [•]")
merge_val("B11:C11", "Attn: [Contact Name]")
merge_val("B12:C12", "Client Email: [name@client.com]")

merge_val("E7:F7", "PROJECT / ENGAGEMENT", size=10, bold=True, color="FFFFFF", fill=NAVY)
merge_val("E8:F8", "[Engagement name, e.g. ISO 9001 Audit Readiness]")
merge_val("E9:F9", "PO / Reference: [•]")
merge_val("E10:F10", "Payment Terms: Net [15] days")

# Line items header
row0 = 14
headers = ["#", "Description", "Qty / Hrs", "Unit Rate (AED)", "Amount (AED)"]
for i, h in enumerate(headers):
    col = ["B", "C", "D", "E", "F"][i]
    cell = ws[f"{col}{row0}"]
    cell.value = h
    cell.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = border

# Example / editable rows
sample_rows = [
    (1, "ISO 9001 Audit Readiness — Gap Assessment (example)", 1, 8000),
    (2, "", "", ""),
    (3, "", "", ""),
    (4, "", "", ""),
    (5, "", "", ""),
]
r = row0 + 1
first_item_row = r
for num, desc, qty, rate in sample_rows:
    ws[f"B{r}"] = num
    ws[f"C{r}"] = desc
    ws[f"D{r}"] = qty
    ws[f"E{r}"] = rate
    ws[f"F{r}"] = f"=IF(AND(D{r}<>\"\",E{r}<>\"\"),D{r}*E{r},\"\")"
    for col in ["B", "C", "D", "E", "F"]:
        c = ws[f"{col}{r}"]
        c.border = border
        c.font = Font(name=FONT, size=10)
        if col in ("D", "E", "F"):
            c.alignment = Alignment(horizontal="right")
        if col == "E":
            c.number_format = '#,##0.00'
        if col == "F":
            c.number_format = '#,##0.00'
    ws[f"C{r}"].fill = PatternFill("solid", fgColor=INPUT_FILL)
    ws[f"D{r}"].fill = PatternFill("solid", fgColor=INPUT_FILL)
    ws[f"E{r}"].fill = PatternFill("solid", fgColor=INPUT_FILL)
    r += 1
last_item_row = r - 1

# Totals
totals_row = r + 1
ws[f"E{totals_row}"] = "Subtotal"
ws[f"F{totals_row}"] = f"=SUM(F{first_item_row}:F{last_item_row})"
ws[f"E{totals_row+1}"] = "VAT (5%)"
ws[f"F{totals_row+1}"] = f"=F{totals_row}*0.05"
ws[f"E{totals_row+2}"] = "TOTAL DUE (AED)"
ws[f"F{totals_row+2}"] = f"=F{totals_row}+F{totals_row+1}"

for rr in (totals_row, totals_row + 1):
    ws[f"E{rr}"].font = Font(name=FONT, size=10, bold=True)
    ws[f"E{rr}"].alignment = Alignment(horizontal="right")
    ws[f"F{rr}"].font = Font(name=FONT, size=10)
    ws[f"F{rr}"].number_format = '#,##0.00'
    ws[f"F{rr}"].alignment = Alignment(horizontal="right")
    ws[f"F{rr}"].border = border

ws[f"E{totals_row+2}"].font = Font(name=FONT, size=11, bold=True, color="FFFFFF")
ws[f"E{totals_row+2}"].fill = PatternFill("solid", fgColor=NAVY)
ws[f"E{totals_row+2}"].alignment = Alignment(horizontal="right")
ws[f"F{totals_row+2}"].font = Font(name=FONT, size=11, bold=True, color="FFFFFF")
ws[f"F{totals_row+2}"].fill = PatternFill("solid", fgColor=NAVY)
ws[f"F{totals_row+2}"].number_format = '#,##0.00'
ws[f"F{totals_row+2}"].alignment = Alignment(horizontal="right")

# Bank details + notes
bank_row = totals_row + 4
merge_val(f"B{bank_row}:C{bank_row}", "BANK DETAILS", size=10, bold=True, color="FFFFFF", fill=NAVY)
merge_val(f"B{bank_row+1}:C{bank_row+1}", "Bank Name: [•]")
merge_val(f"B{bank_row+2}:C{bank_row+2}", "Account Name: Qflow Consulting")
merge_val(f"B{bank_row+3}:C{bank_row+3}", "IBAN: [•]")
merge_val(f"B{bank_row+4}:C{bank_row+4}", "Swift/BIC: [•]")

merge_val(f"E{bank_row}:F{bank_row}", "NOTES", size=10, bold=True, color="FFFFFF", fill=NAVY)
merge_val(f"E{bank_row+1}:F{bank_row+4}", "Thank you for your business. Late payments accrue interest per the signed Consulting Services Agreement, Clause 3.", italic=True)

# Authorized signature / stamp zone — Qflow Tools > Insert Company Stamp drops an image here
stamp_row = bank_row + 6
merge_val(f"B{stamp_row}:F{stamp_row}", "AUTHORIZED SIGNATURE & STAMP", size=10, bold=True, color="FFFFFF", fill=NAVY)
for rr in range(stamp_row + 1, stamp_row + 5):
    for col in ["B", "C", "D", "E", "F"]:
        ws[f"{col}{rr}"].border = border

# Legend
legend_row = stamp_row + 6
merge_val(f"B{legend_row}:F{legend_row}", "Yellow cells = fields to edit. Formulas (Amount, Subtotal, VAT, Total) recalculate automatically — do not overwrite. Use the Qflow Invoice Tools menu to assign the invoice number, generate the PDF, insert the stamp, and email the client.", size=9, italic=True, color="808080")

ws.sheet_view.showGridLines = False

# ---- Config sheet (hidden) ----
cfg = wb.create_sheet("Config")
cfg["A1"] = "Next Invoice Number"
cfg["B1"] = 1
cfg["A2"] = "Invoice Number Prefix"
cfg["B2"] = "INV-"
cfg["A3"] = "Invoices Root Folder ID (leave blank to auto-create next to this sheet)"
cfg["B3"] = ""
for r_ in (1, 2, 3):
    cfg[f"A{r_}"].font = Font(name=FONT, size=10, bold=True)
    cfg[f"B{r_}"].font = Font(name=FONT, size=10)
    cfg.column_dimensions["A"].width = 55
    cfg.column_dimensions["B"].width = 20
cfg.sheet_state = "hidden"

# ---- Invoice Log sheet ----
log = wb.create_sheet("Invoice Log")
log_headers = ["Invoice No.", "Date Generated", "Client", "Total (AED)", "Status", "PDF Link"]
for i, h in enumerate(log_headers, start=1):
    c = log.cell(row=1, column=i, value=h)
    c.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor=NAVY)
    c.alignment = Alignment(horizontal="center", vertical="center")
log.column_dimensions["A"].width = 14
log.column_dimensions["B"].width = 16
log.column_dimensions["C"].width = 28
log.column_dimensions["D"].width = 14
log.column_dimensions["E"].width = 12
log.column_dimensions["F"].width = 40
log.freeze_panes = "A2"

wb.save("Qflow_Invoice_Template.xlsx")
print("written")
